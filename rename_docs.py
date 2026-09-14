import os
try:
    import openpyxl
    from docx import Document
except ImportError:
    import subprocess
    import sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "python-docx"])
    import openpyxl
    from docx import Document

docs_dir = r"d:\Flutter_Project\DTCproduct\docs\1 Máy tách màu"

# 1. Update docx
docx_path = os.path.join(docs_dir, "Dac_diem_thong_so_ky_thuat_may_SC.docx")
if os.path.exists(docx_path):
    doc = Document(docx_path)
    for p in doc.paragraphs:
        if 'SC12 Pro' in p.text:
            p.text = p.text.replace('SC12 Pro', 'SC16 Pro')
        if 'SC12 pro' in p.text:
            p.text = p.text.replace('SC12 pro', 'SC16 Pro')
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if 'SC12 Pro' in cell.text:
                    cell.text = cell.text.replace('SC12 Pro', 'SC16 Pro')
                if 'SC12 pro' in cell.text:
                    cell.text = cell.text.replace('SC12 pro', 'SC16 Pro')
    doc.save(docx_path)
    print(f"Updated {docx_path}")

# 2. Update xlsx
xlsx_files = [
    "1.1 Thong_So_MayTachMauSC.xlsx",
    "1.3 Phan_Tich_Hoan_Von_MTM.xlsx",
    "1.2 Thiet_Bi_Phu_Tro_MTM.xlsx",
    "1.5 Tra_Cuu_Loi_May_Tach_Mau.xlsx"
]

for filename in xlsx_files:
    filepath = os.path.join(docs_dir, filename)
    if os.path.exists(filepath):
        wb = openpyxl.load_workbook(filepath)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        if 'SC12 Pro' in cell.value:
                            cell.value = cell.value.replace('SC12 Pro', 'SC16 Pro')
                        if 'SC12 pro' in cell.value:
                            cell.value = cell.value.replace('SC12 pro', 'SC16 Pro')
        wb.save(filepath)
        print(f"Updated {filepath}")
