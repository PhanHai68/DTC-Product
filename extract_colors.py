from docx import Document
import os

doc_path = r"D:\Flutter_Project\DTCproduct\docs\1 Máy tách màu\Dac_diem_thong_so_ky_thuat_may_SC.docx"
doc = Document(doc_path)

found = False
for i, p in enumerate(doc.paragraphs):
    if "ỨNG DỤNG" in p.text.upper():
        print(f"--- FOUND SECTION: {p.text} ---")
        found = True
        for j in range(i, min(i+20, len(doc.paragraphs))):
            para = doc.paragraphs[j]
            if not para.text.strip(): continue
            print(f"\n[Paragraph]")
            for run in para.runs:
                if not run.text.strip(): continue
                color = run.font.color.rgb if run.font.color and run.font.color.rgb else "None"
                print(f"  Run: {run.text.strip()} | Color: {color}")
        break
