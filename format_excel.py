import openpyxl
import re

file_path = r"D:\Flutter_Project\DTCproduct\docs\2 Máy nén khí\CAC_LOI_THUONG_GAP_CUA_MNK.xlsx"
wb = openpyxl.load_workbook(file_path)

def format_text(text):
    if not text:
        return text
    text = str(text).strip()
    if not text:
        return text
    parts = re.split(r'[/_\n]', text)
    formatted = []
    for p in parts:
        p = p.strip()
        if not p:
            continue
        if p.startswith('-'): p = p[1:].strip()
        if p.startswith('+'): p = p[1:].strip()
        if p.startswith('_'): p = p[1:].strip()
        
        p = re.sub(r'[;:,.\s]+$', '', p)
        if p:
            p = p[0].upper() + p[1:]
            formatted.append(f'- {p}.')
    return '\n'.join(formatted)

def format_title(text):
    if not text:
        return text
    text = str(text).strip()
    text = re.sub(r'[;:,.\s]+$', '', text)
    if text:
        text = text[0].upper() + text[1:]
    return text

# Sheet 1: Lỗi_MNK_Thường_gặp
if 'Lỗi_MNK_Thường_gặp' in wb.sheetnames:
    ws = wb['Lỗi_MNK_Thường_gặp']
    for row in range(4, ws.max_row + 1):
        if ws.cell(row=row, column=2).value:
            ws.cell(row=row, column=2).value = format_title(ws.cell(row=row, column=2).value)
        if ws.cell(row=row, column=3).value:
            ws.cell(row=row, column=3).value = format_text(ws.cell(row=row, column=3).value)

# Sheet 2: Lỗi_MNK_HDSD_BĐK
if 'Lỗi_MNK_HDSD_BĐK' in wb.sheetnames:
    ws = wb['Lỗi_MNK_HDSD_BĐK']
    for row in range(4, ws.max_row + 1):
        if ws.cell(row=row, column=2).value:
            ws.cell(row=row, column=2).value = format_title(ws.cell(row=row, column=2).value)
        if ws.cell(row=row, column=3).value:
            ws.cell(row=row, column=3).value = format_text(ws.cell(row=row, column=3).value)
        if ws.cell(row=row, column=4).value:
            ws.cell(row=row, column=4).value = format_text(ws.cell(row=row, column=4).value)

# Sheet 3: Lỗi_MNK_Hao_Dầu
if 'Lỗi_MNK_Hao_Dầu' in wb.sheetnames:
    ws = wb['Lỗi_MNK_Hao_Dầu']
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=2).value:
            ws.cell(row=row, column=2).value = format_title(ws.cell(row=row, column=2).value)
        if ws.cell(row=row, column=3).value:
            ws.cell(row=row, column=3).value = format_text(ws.cell(row=row, column=3).value)
        if ws.cell(row=row, column=4).value:
            ws.cell(row=row, column=4).value = format_text(ws.cell(row=row, column=4).value)

wb.save(file_path)
print("Excel updated")
